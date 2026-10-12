"""Generic OOXML preserves characterized payloads, original bytes and validation order."""

import base64
import io
import json
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook

from report_generator.v2 import (
    MetadataFormat,
    annotate_workbook,
    original_workbook,
    package_fingerprint,
    read_metadata,
)
from report_generator.v2.workbook import _parse

FORMAT = MetadataFormat(
    'customXml/clearml-yolo-identity.xml',
    'https://clearml-yolo/model-identity/v1',
    'identity',
    '1',
    'clearmlYoloIdentity',
)
FIXTURE = Path(__file__).parent / 'fixtures/ooxml-characterization.json'


def characterized():
    value = json.loads(FIXTURE.read_text())
    return (
        base64.b64decode(value['original']),
        {key: base64.b64decode(content) for key, content in value['parts'].items()},
        value,
    )


def package(parts):
    output = io.BytesIO()
    with ZipFile(output, 'w', compression=ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    return output.getvalue()


def test_characterized_v1_payloads_byte_exact_recovery_and_reannotation(tmp_path):
    original, expected, fixture = characterized()
    path = tmp_path / 'book.xlsx'
    path.write_bytes(original)
    annotate_workbook(
        path,
        banner_lines=fixture['banner_lines'],
        metadata_texts=fixture['metadata_texts'],
        format=FORMAT,
    )
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    assert parts == expected
    assert original_workbook(path.read_bytes(), FORMAT) == original
    assert original_workbook(original, FORMAT) == original
    assert read_metadata(original, FORMAT) is None
    first = parts.copy()
    annotate_workbook(
        path,
        banner_lines=fixture['banner_lines'],
        metadata_texts=fixture['metadata_texts'],
        format=FORMAT,
    )
    with ZipFile(path) as archive:
        assert {name: archive.read(name) for name in archive.namelist()} == first
    workbook = load_workbook(path)
    shift = len(fixture['banner_lines'])
    main = workbook['Main sheet']
    other = workbook['Other']
    assert main.cell(3 + shift, 2).value == f'=SUM($B${2 + shift}:B{2 + shift})'
    assert other.cell(1 + shift, 1).value == f"='Main sheet'!$B${2 + shift}+1"
    assert main.tables['Results'].ref == f'A{1 + shift}:B{2 + shift}'
    assert main.auto_filter.ref == f'A{1 + shift}:B{2 + shift}'
    assert main.data_validations.dataValidation[0].formula1 == f'B{2 + shift}'
    assert main.freeze_panes == f'B{2 + shift}'
    assert main.row_breaks.brk[0].id == 2 + shift
    assert main.print_title_cols == '$A:$A'
    for sheet in workbook:
        assert sheet.page_setup.fitToWidth == 1
        assert sheet.print_title_rows.startswith('$1:')
        assert all(sheet.cell(row, 1).data_type == 's' for row in range(1, shift + 1))
    workbook.close()


@pytest.mark.parametrize('part', ['xl/worksheets/sheet1.xml', FORMAT.path, 'xl/styles.xml'])
def test_fingerprint_rejects_edited_annotated_parts(part):
    _, parts, _ = characterized()
    if part == FORMAT.path:
        root = ET.fromstring(parts[part])
        root.find(f'{{{FORMAT.namespace}}}models').text = '{}'
        parts[part] = ET.tostring(root, encoding='utf-8', xml_declaration=True)
    else:
        parts[part] = parts[part].replace(b'>', b'> ', 1)
    with pytest.raises(ValueError, match='changed after'):
        original_workbook(package(parts), FORMAT)


def test_schema_owner_validator_runs_before_fingerprint_comparison():
    _, parts, _ = characterized()
    root = ET.fromstring(parts[FORMAT.path])
    root.find(f'{{{FORMAT.namespace}}}models').text = 'invalid'
    parts[FORMAT.path] = ET.tostring(root, encoding='utf-8', xml_declaration=True)

    def validator(metadata):
        raise ValueError('schema owner rejected data')

    with pytest.raises(ValueError, match='schema owner rejected'):
        original_workbook(package(parts), FORMAT, validator)


@pytest.mark.parametrize(
    'change,error',
    [('version', 'Unsupported'), ('fingerprint', 'fingerprint'), ('original', 'original content')],
)
def test_metadata_required_envelope_fields(change, error):
    _, parts, _ = characterized()
    root = ET.fromstring(parts[FORMAT.path])
    if change == 'version':
        root.set('version', '2')
    elif change == 'fingerprint':
        root.attrib.pop('package_sha256')
    else:
        root.remove(root.find(f'{{{FORMAT.namespace}}}original'))
    parts[FORMAT.path] = ET.tostring(root, encoding='utf-8', xml_declaration=True)
    if change == 'original':
        root.set('package_sha256', package_fingerprint(parts, FORMAT))
        parts[FORMAT.path] = ET.tostring(root, encoding='utf-8', xml_declaration=True)
    with pytest.raises(ValueError, match=error):
        original_workbook(package(parts), FORMAT)


@pytest.mark.parametrize('content', [b'<!DOCTYPE x><x/>', b'<!ENTITY x "y"><x/>', b'\x00<x/>'])
def test_unsafe_xml_rejected(content):
    with pytest.raises(ValueError, match='DTDs or entities'):
        _parse(content)


def test_nonapplication_metadata_envelope_and_literal_banner(tmp_path):
    original, _, _ = characterized()
    path = tmp_path / 'book.xlsx'
    path.write_bytes(original)
    format = MetadataFormat(
        'customXml/acme.xml', 'urn:acme:metadata', 'report', '2026', 'acmeReport'
    )
    annotate_workbook(
        path,
        banner_lines=['=literal', '  Unicode 模型  '],
        metadata_texts={'owner': 'Acme', 'run': 'full-id'},
        format=format,
    )
    metadata = read_metadata(path.read_bytes(), format)
    assert metadata.findtext('{urn:acme:metadata}owner') == 'Acme'
    assert original_workbook(path.read_bytes(), format) == original
    workbook = load_workbook(path)
    for sheet in workbook:
        assert sheet['A1'].value == '=literal' and sheet['A1'].data_type == 's'
        assert sheet['A2'].value == '  Unicode 模型  '
    workbook.close()
