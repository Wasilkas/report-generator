import importlib.util
import math
import random
from contextlib import nullcontext
from dataclasses import replace
from pathlib import Path

import pandas as pd
import pytest
from click.testing import CliRunner
from openpyxl import Workbook, load_workbook

from report_generator.app_config import CLASS_COL
from report_generator.cli import _coerce_sheet, cli
from report_generator.config import Config
from report_generator.core.calculator import ClassFilter, ComparisonCalculator, MeanRowCalculator
from report_generator.core.reader import MetricsReader
from report_generator.reports.business import utils
from report_generator.reports.business.builder import BusinessReportBuilder
from report_generator.reports.dev.builder import DevReportBuilder


@pytest.fixture
def config():
    return Config.default()


def inputs(tmp_path, **columns):
    data = {
        CLASS_COL: ['cat'],
        'Количество примеров train': [20],
        'f1_score': [0.8],
        'perebrak': [0.3],
        'nedobrak': [0.2],
    }
    data.update(columns)
    paths = [tmp_path / 'new.xlsx', tmp_path / 'prod.xlsx']
    for path in paths:
        pd.DataFrame(data).to_excel(path, index=False)
    return paths


def test_reader_normalization(tmp_path):
    path = tmp_path / 'raw.xlsx'
    pd.DataFrame({'': [' cat ', '0.50', '1e-3', 'cat2', ' Среднее ', ' ']}).to_excel(
        path, index=False
    )
    assert MetricsReader(path).read()[CLASS_COL].tolist() == ['cat', 'cat2']
    pd.DataFrame({'': ['cat', ' cat ']}).to_excel(path, index=False)
    with pytest.raises(ValueError, match='Duplicate'):
        MetricsReader(path).read()


@pytest.mark.parametrize('count', ['bad', None, math.inf, -1, 1.5, True])
def test_invalid_training_count(config, count):
    with pytest.raises(ValueError, match='training'):
        ClassFilter(config).split(pd.DataFrame({CLASS_COL: ['a'], 'Число примеров train': [count]}))


def test_training_schema_cutoff(config):
    frame = pd.DataFrame({CLASS_COL: ['a', 'b'], 'Число примеров train': [19, 20]})
    included, excluded = ClassFilter(config).split(frame)
    assert included[CLASS_COL].tolist() == ['b']
    assert '< 20' in excluded.iloc[0]['Причина']
    with pytest.warns(UserWarning, match='training'):
        ClassFilter(config).split(frame.drop(columns=['Число примеров train']))
    with pytest.raises(ValueError, match='Multiple'):
        ClassFilter(config).split(frame.assign(**{'Количество примеров train': [20, 20]}))


def test_macro_and_numeric_differences(config):
    a = pd.DataFrame(
        {
            CLASS_COL: ['a', 'b'],
            'f1_score': [0.9, 0.3],
            'confidence': [0.5, 0.5],
            'tp': [20, 100],
            'text': ['a', 'b'],
        }
    )
    assert MeanRowCalculator(config).append_mean_row(a).iloc[-1]['f1_score'] == pytest.approx(0.6)
    b = a.copy()
    b['f1_score'] = ['bad', None]
    diff = ComparisonCalculator(config).compute(a, b)
    assert diff.columns.tolist() == [CLASS_COL, 'f1_score']
    assert diff['f1_score'].isna().all()


@pytest.mark.parametrize('builder', [DevReportBuilder, BusinessReportBuilder])
def test_workbook_provenance_id_and_output(tmp_path, config, builder):
    paths = inputs(tmp_path, **{'ID': ['external']})
    prod = pd.read_excel(paths[1])
    prod['Количество примеров train'] = 5
    prod.to_excel(paths[1], index=False)
    output = tmp_path / 'report.xlsx'
    instance = builder(*(MetricsReader(p) for p in paths), config)
    instance.build(output)
    wb = load_workbook(output)
    excluded = wb[config.sheet_names.excluded]
    assert excluded['A2'].value == CLASS_COL
    rows = list(excluded.iter_rows(min_row=3, values_only=True))
    assert any(5 in row and 'Прод модель' in row for row in rows)
    assert not any('Только' in str(value) for row in rows for value in row)
    assert '≤' not in excluded['A1'].value
    with pytest.raises(FileExistsError):
        instance.build(output)
    instance.build(output, overwrite=True)
    for source in paths:
        with pytest.raises(ValueError, match='input'):
            instance.build(source, overwrite=True)
    alias = tmp_path / 'alias.xlsx'
    alias.symlink_to(paths[0])
    with pytest.raises(ValueError, match='input'):
        instance.build(alias, overwrite=True)
    alias.unlink()
    alias.hardlink_to(paths[1])
    with pytest.raises(ValueError, match='input'):
        instance.build(alias, overwrite=True)


def test_id_preserved(tmp_path, config):
    paths = inputs(tmp_path, ID=['custom'])
    output = tmp_path / 'out.xlsx'
    DevReportBuilder(*(MetricsReader(p) for p in paths), config).build(output)
    sheet = load_workbook(output)[config.sheet_names.model1]
    column = next(cell.column for cell in sheet[2] if cell.value == 'ID')
    assert sheet.cell(3, column).value == 'custom'


def test_business_boundaries_invalid_and_gross(config):
    biz = config.business
    df = utils.add_goal_columns(
        pd.DataFrame({'perebrak': [0.3, 0.6, 0.7, None], 'nedobrak': [0.2, 0.4, 0.1, 0.1]}), biz
    )
    assert df[biz.col_perebrak_target].tolist()[:3] == [1, 0, 0]
    assert df[biz.col_perebrak_gross].tolist()[:3] == [0, 0, 1]
    assert pd.isna(df.iloc[-1][biz.col_perebrak_target])
    assert utils.classes_in_target(df, biz) is None
    assert utils.classes_gross_not_target(df.iloc[:3], biz) == pytest.approx(1 / 3)


@pytest.mark.parametrize('new,prod,expected', [(0, 0, [1, 1, 1]), (0.8, 0, [2, 1, 1])])
def test_verdict_zero_baselines(config, new, prod, expected):
    def frame(f1):
        return utils.add_goal_columns(
            pd.DataFrame({'f1_score': [f1], 'perebrak': [0.1], 'nedobrak': [0.1]}), config.business
        )

    sheet = Workbook().active
    utils.write_verdict_sheet(sheet, frame(new), frame(prod), config)
    assert [sheet.cell(r, 5).value for r in range(3, 6)] == expected
    assert sheet['D3'].value == 'NA'
    assert any('без изменений' in str(c.value).lower() for row in sheet for c in row)
    utils.write_verdict_sheet(sheet, frame(new).drop(columns=['f1_score']), frame(prod), config)
    assert sheet['E6'].value == 'Недостаточно данных'


def test_custom_translations(config):
    biz = replace(
        config.business,
        column_translations={
            'f1_score': 'F1 custom',
            'perebrak': 'P custom',
            config.business.col_perebrak_target: 'Goal custom',
        },
    )
    config = replace(config, business=biz)
    assert utils.build_comparison_color_fn(config)('F1 custom', 0.1) is not None
    assert utils.build_comparison_color_fn(config)('P custom', 6) is not None
    assert utils.build_format_map(config)['Goal custom'] == '0.00'
    with pytest.raises(ValueError, match='collision'):
        utils.translate_columns(pd.DataFrame({'a': [1], 'b': [2]}), {'a': 'b'})


@pytest.mark.parametrize(
    'yaml',
    [
        'unknown: 1',
        'business:\n  unknown: 1',
        'min_train_count: true',
        'degradation_threshold: -1',
        'business: []',
        'colors: []',
        '[]',
        'business:\n  target_perebrak: .nan',
        'sheet_names:\n  model1: Same\n  model2: same',
        'sheet_names:\n  verdict: "bad/name"',
        'ratio_cols: precision',
        'business:\n  target_nedobrak: 2',
    ],
)
def test_config_rejects_invalid(tmp_path, yaml):
    path = tmp_path / 'config.yaml'
    path.write_text(yaml)
    with pytest.raises(ValueError):
        Config.load(path)


def test_cli_selectors_errors_and_force(tmp_path):
    assert _coerce_sheet('name:1') == '1'
    assert _coerce_sheet('index:1') == 1
    assert _coerce_sheet('1') == 1
    with pytest.raises(ValueError):
        _coerce_sheet('index:-1')
    paths = inputs(tmp_path)
    runner = CliRunner()
    output = tmp_path / 'out.xlsx'
    args = ['dev', *map(str, paths), str(output)]
    assert runner.invoke(cli, args).exit_code == 0
    assert runner.invoke(cli, args).exit_code != 0
    assert runner.invoke(cli, args + ['--force']).exit_code == 0
    paths[0].write_text('invalid workbook')
    result = runner.invoke(cli, args + ['--force'])
    assert result.exit_code != 0 and 'Error:' in result.output


def test_sample_import_and_generation(tmp_path):
    state = random.getstate()
    spec = importlib.util.spec_from_file_location(
        'sample', Path(__file__).parents[1] / 'create_test_data.py'
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    assert random.getstate() == state
    paths = [tmp_path / 'a.xlsx', tmp_path / 'b.xlsx']
    module.generate(*paths)
    before = pd.read_excel(paths[0])
    with pytest.raises(FileExistsError):
        module.generate(*paths)
    module.generate(*paths, overwrite=True)
    pd.testing.assert_frame_equal(before, pd.read_excel(paths[0]))


def test_finite_macro_coverage_and_comparison_identity(config):
    frame = pd.DataFrame(
        {
            CLASS_COL: ['a', 'b'],
            'ap50': [0.5, math.inf],
            'Unfamiliar count': [4, 8],
            'f1_score': [0.5, 0.7],
        }
    )
    row = MeanRowCalculator(config).append_mean_row(frame).iloc[-1]
    assert row['ap50'] == 0.5
    assert row['ap50 coverage'] == '1/2'
    changed = replace(config, excluded_from_mean=frozenset({'f1_score'}))
    diff = ComparisonCalculator(changed).compute(frame, frame)
    assert 'f1_score' in diff and 'Unfamiliar count' not in diff


@pytest.mark.parametrize(
    'yaml',
    [
        'null',
        'colors: null',
        'sheet_names: null',
        'business: null',
        'business:\n  column_translations: null',
        'better_lower_cols: [f1_score]',
        'business:\n  col_perebrak_target: f1_score',
        'business:\n  col_perebrak_target: Same\n  col_nedobrak_target: Same',
    ],
)
def test_config_null_identity_conflicts(tmp_path, yaml):
    path = tmp_path / 'bad.yaml'
    path.write_text(yaml)
    with pytest.raises(ValueError):
        Config.load(path)


def test_business_unavailable_warns_and_equality_zero(config):
    with pytest.warns(UserWarning, match='Unavailable'):
        frame = utils.add_goal_columns(
            pd.DataFrame({'f1_score': [True], 'perebrak': [math.inf], 'nedobrak': ['bad']}),
            config.business,
        )
    assert pd.isna(frame['f1_score'].iloc[0])
    assert pd.isna(frame['perebrak'].iloc[0])
    assert utils.score_higher_better(0, 0) == 1
    assert utils.score_lower_better(0, 0) == 1


def test_translated_class_mean_saved_styles(tmp_path, config):
    biz = replace(
        config.business,
        column_translations={
            CLASS_COL: 'Category',
            'f1_score': 'Quality',
            'perebrak': 'P',
            'nedobrak': 'N',
        },
    )
    custom = replace(config, business=biz)
    paths = inputs(tmp_path)
    output = tmp_path / 'custom.xlsx'
    BusinessReportBuilder(*(MetricsReader(p) for p in paths), custom).build(output)
    ws = load_workbook(output)[custom.sheet_names.model1]
    assert ws['A2'].value == 'Category'
    assert ws['A4'].font.italic
    assert ws['A4'].fill.fgColor.rgb.endswith(config.colors.mean)
    assert ws['A3'].alignment.horizontal == 'left'


def test_sample_preflight_prevents_partial_write(tmp_path):
    spec = importlib.util.spec_from_file_location(
        'sample', Path(__file__).parents[1] / 'create_test_data.py'
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    good = tmp_path / 'new.xlsx'
    bad = tmp_path / 'directory.xlsx'
    bad.mkdir()
    with pytest.raises((ValueError, OSError)):
        module.generate(good, bad, overwrite=True)
    assert not good.exists()


@pytest.mark.parametrize('mode', ['dev', 'business'])
def test_saved_workbook_unavailable_and_legends(tmp_path, config, mode):
    paths = inputs(tmp_path, ap50=[float('nan')])
    prod = pd.read_excel(paths[1])
    prod['f1_score'] = 'bad'
    prod.to_excel(paths[1], index=False)
    output = tmp_path / 'out.xlsx'
    runner = CliRunner()
    with pytest.warns(UserWarning, match='Unavailable') if mode == 'business' else nullcontext():
        result = runner.invoke(cli, [mode, *map(str, paths), str(output)])
    assert result.exit_code == 0, result.output
    wb = load_workbook(output)
    ws = wb[config.sheet_names.comparison]
    f1 = next(c.column for c in ws[2] if c.value == 'f1_score')
    assert ws.cell(3, f1).value == 'NA'
    assert ws.cell(3, f1).fill.patternType is None
    if mode == 'business':
        assert wb[config.sheet_names.verdict]['E6'].value == 'Недостаточно данных'
        assert any('5.0 п.п.' in str(c.value) for row in ws for c in row)
    first = wb[config.sheet_names.model1]
    ap = next(c.column for c in first[2] if c.value == 'ap50')
    assert first.cell(3, ap).value == 'NA'
    coverage = next(c.column for c in first[2] if c.value == 'ap50 coverage')
    assert first.cell(4, coverage).value == '0/1'


@pytest.mark.parametrize('mode', ['dev', 'business'])
def test_numeric_sheet_name_and_custom_verdict(tmp_path, mode):
    paths = inputs(tmp_path)
    for path in paths:
        data = pd.read_excel(path)
        with pd.ExcelWriter(path) as writer:
            data.assign(**{CLASS_COL: ['wrong']}).to_excel(writer, sheet_name='first', index=False)
            data.to_excel(writer, sheet_name='1', index=False)
    config = tmp_path / 'cfg.yaml'
    config.write_text('sheet_names:\n  verdict: Decision\n', encoding='utf-8')
    output = tmp_path / 'out.xlsx'
    result = CliRunner().invoke(
        cli,
        [
            '--config',
            str(config),
            mode,
            *map(str, paths),
            str(output),
            '--sheet1',
            'name:1',
            '--sheet2',
            'name:1',
        ],
    )
    assert result.exit_code == 0, result.output
    wb = load_workbook(output)
    assert wb['Новая модель']['A3'].value == 'cat'
    if mode == 'business':
        assert 'Decision' in wb.sheetnames


@pytest.mark.parametrize('ids', [[None], [' '], [True], [math.inf]])
def test_invalid_ids(tmp_path, ids):
    path = inputs(tmp_path, ID=ids)[0]
    with pytest.raises(ValueError, match='ID'):
        MetricsReader(path).read()


def test_rollout_partial_class_and_zero_gross(config):
    biz = config.business
    with pytest.warns(UserWarning):
        new = utils.add_goal_columns(
            pd.DataFrame({'f1_score': [0.9, None], 'perebrak': [0.1, 0.1], 'nedobrak': [0.1, 0.1]}),
            biz,
        )
    prod = utils.add_goal_columns(
        pd.DataFrame({'f1_score': [0.8, 0.8], 'perebrak': [0.1, 0.1], 'nedobrak': [0.1, 0.1]}), biz
    )
    assert MeanRowCalculator(config).append_mean_row(new).iloc[-1]['f1_score'] == 0.9
    ws = Workbook().active
    utils.write_verdict_sheet(ws, new, prod, config)
    assert ws['E6'].value == 'Недостаточно данных'
    new = utils.add_goal_columns(
        pd.DataFrame({'f1_score': [0.9], 'perebrak': [0.7], 'nedobrak': [0.1]}), biz
    )
    utils.write_verdict_sheet(ws, new, prod.iloc[:1], config)
    assert ws['E5'].value == 0 and ws['D5'].value == 'NA'
    assert ws['E6'].value == 'Не к выкатке'


def test_cli_unexpected_internal_error_is_not_masked(tmp_path, monkeypatch):
    paths = inputs(tmp_path)

    def fail(*args, **kwargs):
        raise RuntimeError('internal defect')

    monkeypatch.setattr(DevReportBuilder, 'build', fail)
    result = CliRunner().invoke(cli, ['dev', *map(str, paths), str(tmp_path / 'out.xlsx')])
    assert isinstance(result.exception, RuntimeError)
    assert 'Cannot generate report' not in result.output


@pytest.mark.parametrize('builder', [DevReportBuilder, BusinessReportBuilder])
def test_class_formula_like_names_remain_literal(tmp_path, config, builder):
    paths = inputs(tmp_path, **{CLASS_COL: ['=literal']})
    # Excel input must carry a literal string, too.
    for path in paths:
        wb = load_workbook(path)
        wb.active['A2'].data_type = 's'
        wb.save(path)
    output = tmp_path / 'out.xlsx'
    builder(*(MetricsReader(p) for p in paths), config).build(output)
    wb = load_workbook(output)
    assert wb[config.sheet_names.model1]['A3'].data_type == 's'
    prod = load_workbook(paths[1])
    prod.active['B2'] = 5
    prod.save(paths[1])
    builder(*(MetricsReader(p) for p in paths), config).build(output, overwrite=True)
    excluded = load_workbook(output)[config.sheet_names.excluded]
    assert excluded['A3'].value == '=literal' and excluded['A3'].data_type == 's'


def test_recursive_config_keys_have_actionable_cli_error(tmp_path):
    path = tmp_path / 'bad.yaml'
    path.write_text('colors:\n  1: FF0000\n  nope: FF0000\n')
    with pytest.raises(ValueError, match='keys must be strings'):
        Config.load(path)
    paths = inputs(tmp_path)
    result = CliRunner().invoke(
        cli, ['--config', str(path), 'business', *map(str, paths), str(tmp_path / 'out.xlsx')]
    )
    assert result.exit_code != 0
    assert 'Error: Invalid configuration' in result.output
    assert 'keys must be strings' in result.output


@pytest.mark.parametrize(
    'source', ['Класс', 'ID', 'Количество примеров train', 'unfamiliar_count', 'confidence', 'tp']
)
@pytest.mark.parametrize('key', ['f1_col', 'perebrak_col', 'nedobrak_col'])
def test_business_source_alias_cannot_mutate_metadata(tmp_path, source, key):
    path = tmp_path / 'bad.yaml'
    path.write_text(f'business:\n  {key}: "{source}"\n', encoding='utf-8')
    with pytest.raises(ValueError, match='metadata'):
        Config.load(path)


def test_business_legitimate_metric_source_alias(tmp_path):
    path = tmp_path / 'valid.yaml'
    path.write_text('business:\n  f1_col: precision\n  perebrak_col: custom_error\n')
    config = Config.load(path)
    assert config.business.f1_col == 'precision'
    assert config.business.perebrak_col == 'custom_error'


@pytest.mark.parametrize('reverse', [False, True])
def test_sample_rejects_nested_destinations_before_write(tmp_path, reverse):
    spec = importlib.util.spec_from_file_location(
        'sample', Path(__file__).parents[1] / 'create_test_data.py'
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    paths = [tmp_path / 'first.xlsx', tmp_path / 'first.xlsx' / 'second.xlsx']
    if reverse:
        paths.reverse()
    with pytest.raises(ValueError, match='nested'):
        module.generate(*paths)
    assert not (tmp_path / 'first.xlsx').exists()


@pytest.mark.parametrize(
    'override',
    [
        'business:\n  perebrak_col: precision',
        'business:\n  nedobrak_col: recall',
        'business:\n  f1_col: perebrak',
        'better_higher_cols: [custom_error]\nbusiness:\n  perebrak_col: custom_error',
        'better_lower_cols: [custom_quality]\nbusiness:\n  f1_col: custom_quality',
    ],
)
def test_business_source_direction_conflicts_rejected(tmp_path, override):
    path = tmp_path / 'direction.yaml'
    path.write_text(override)
    with pytest.raises(ValueError, match='direction'):
        Config.load(path)


def test_custom_business_error_alias_saved_red(tmp_path):
    path = tmp_path / 'alias.yaml'
    path.write_text('business:\n  perebrak_col: custom_error\n  f1_col: custom_quality\n')
    config = Config.load(path)
    paths = inputs(tmp_path, custom_error=[0.5], custom_quality=[0.8])
    prod = pd.read_excel(paths[1])
    prod['custom_error'] = 0.1
    prod['custom_quality'] = 0.7
    prod.to_excel(paths[1], index=False)
    output = tmp_path / 'out.xlsx'
    BusinessReportBuilder(*(MetricsReader(p) for p in paths), config).build(output)
    ws = load_workbook(output)[config.sheet_names.comparison]
    column = next(c.column for c in ws[2] if c.value == 'custom_error')
    assert ws.cell(3, column).value == 40
    assert ws.cell(3, column).fill.fgColor.rgb.endswith(config.colors.negative)
    quality = next(c.column for c in ws[2] if c.value == 'custom_quality')
    assert ws.cell(3, quality).fill.fgColor.rgb.endswith(config.colors.positive)
