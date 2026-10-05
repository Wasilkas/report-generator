"""Reopen reports to verify union display and independent aggregate populations."""

import pandas as pd
import pytest
from openpyxl import load_workbook

from report_generator.app_config import CLASS_COL
from report_generator.config import Config
from report_generator.core.reader import MetricsReader
from report_generator.reports.business.builder import BusinessReportBuilder
from report_generator.reports.dev.builder import DevReportBuilder


def frame(names, scores, counts=None):
    return pd.DataFrame(
        {
            CLASS_COL: names,
            'ID': [f'id-{name}' for name in names],
            'Количество примеров train': counts if counts is not None else [20] * len(names),
            'f1_score': scores,
            'ap50': scores,
            'perebrak': [0.1] * len(names),
            'nedobrak': [0.1] * len(names),
            'confidence': [0.5] * len(names),
            'tp': [0] * len(names),
        }
    )


def table(sheet):
    headers = [c.value for c in sheet[2]]
    rows = {}
    for row in sheet.iter_rows(min_row=3):
        name = row[0].value
        if name == 'Среднее' or name in {'shared', 'new', 'deleted', 'filtered'}:
            rows[name] = {key: cell for key, cell in zip(headers, row) if key}
    return rows


@pytest.mark.parametrize('builder', [DevReportBuilder, BusinessReportBuilder])
@pytest.mark.parametrize('scenario', ['shared', 'disjoint', 'asymmetric'])
def test_union_saved_reports(tmp_path, builder, scenario):
    config = Config.default()
    if scenario == 'disjoint':
        new, prod = frame(['new'], [0]), frame(['deleted'], [0.8])
        names, new_mean, prod_mean, shared_count = ['deleted', 'new'], 0, 0.8, 0
    elif scenario == 'asymmetric':
        new = frame(['shared', 'filtered', 'new'], [0.9, 0.1, 0], [20, 5, 20])
        prod = frame(['shared', 'filtered', 'deleted'], [0.5, 0.7, 0.8])
        names, new_mean, prod_mean, shared_count = (
            ['deleted', 'filtered', 'new', 'shared'],
            0.45,
            2 / 3,
            1,
        )
    else:
        new, prod = frame(['shared', 'new'], [0.9, 0]), frame(['shared', 'deleted'], [0.5, 0.8])
        names, new_mean, prod_mean, shared_count = ['deleted', 'new', 'shared'], 0.45, 0.65, 1
    paths = [tmp_path / 'new.xlsx', tmp_path / 'prod.xlsx']
    for data, path in zip([new, prod], paths):
        data.to_excel(path, index=False)
    output = tmp_path / 'out.xlsx'
    builder(*(MetricsReader(path) for path in paths), config).build(output)
    wb = load_workbook(output)
    one, two, diff = [
        table(wb[name])
        for name in [
            config.sheet_names.model1,
            config.sheet_names.model2,
            config.sheet_names.comparison,
        ]
    ]
    assert list(one) == list(two) == list(diff) == names + ['Среднее']
    assert one['new']['f1_score'].value == 0
    assert one['new']['tp'].value == 0
    for col in ['f1_score', 'ap50', 'ID', 'confidence', 'tp', 'Количество примеров train']:
        display_col = (
            config.business.column_translations.get(col, col)
            if builder is BusinessReportBuilder
            else col
        )
        assert one['deleted'][display_col].value == 'NA'
        assert two['new'][display_col].value == 'NA'
    assert one['new']['ID'].value == 'id-new'
    assert two['deleted']['ID'].value == 'id-deleted'
    for name in ['new', 'deleted']:
        assert diff[name]['f1_score'].value == 'NA'
        assert diff[name]['f1_score'].fill.patternType is None
    assert one['Среднее']['f1_score'].value == pytest.approx(new_mean)
    assert two['Среднее']['f1_score'].value == pytest.approx(prod_mean)
    assert (
        one['Среднее']['f1_score coverage'].value
        == f'{len(new) - (scenario == "asymmetric")}/{len(new) - (scenario == "asymmetric")}'
    )
    if shared_count:
        assert diff['shared']['f1_score'].value == pytest.approx(0.4)
        assert diff['shared']['f1_score'].fill.fgColor.rgb.endswith(config.colors.positive)
        assert diff['Среднее']['f1_score coverage'].value == '1/1'
    else:
        assert diff['Среднее']['f1_score'].value == 'NA'
        assert diff['Среднее']['f1_score coverage'].value == '0/0'
    text = ' '.join(str(c.value) for ws in wb for row in ws for c in row)
    assert 'собственный набор' in text and 'общих классов' in text
    if scenario == 'asymmetric':
        assert one['filtered']['f1_score'].value == 'NA'
        assert two['filtered']['f1_score'].value == 0.7
        assert diff['filtered']['f1_score'].value == 'NA'
        exclusions = ' '.join(str(c.value) for row in wb[config.sheet_names.excluded] for c in row)
        assert 'Мало примеров train' in exclusions and 'другой модели' in exclusions
    if builder is BusinessReportBuilder:
        goal = config.business.col_perebrak_target
        assert one['deleted'][goal].value == 'NA'
        assert two['new'][goal].value == 'NA'
        verdict = wb[config.sheet_names.verdict]
        assert verdict['B3'].value == pytest.approx(new_mean)
        assert verdict['C3'].value == pytest.approx(prod_mean)
        assert verdict['B4'].value == verdict['C4'].value == 1


@pytest.mark.parametrize('builder', [DevReportBuilder, BusinessReportBuilder])
def test_partial_columns_missing_and_per_model_identity(tmp_path, builder):
    config = Config.default()
    new = frame(['shared', 'new'], [0.9, 0])
    new.loc[0, 'ap50'] = float('nan')
    prod = frame(['shared', 'deleted'], [0.5, 0.8]).drop(columns='ap50')
    prod['ID'] = ['prod-shared', 'prod-deleted']
    prod['precision'] = [0.4, 0.6]
    paths = [tmp_path / 'new.xlsx', tmp_path / 'prod.xlsx']
    for data, path in zip([new, prod], paths):
        data.to_excel(path, index=False)
    output = tmp_path / 'out.xlsx'
    instance = builder(*(MetricsReader(path) for path in paths), config)
    eligible1, eligible2, _ = instance._prepare(new, prod)
    assert pd.isna(eligible1.loc[0, 'ap50'])
    assert eligible2['ap50'].isna().all()
    differences = instance._cmp_calc.compute(eligible1, eligible2)
    assert differences['ap50'].isna().all()
    instance.build(output)
    wb = load_workbook(output)
    one, two, diff = [
        table(wb[name])
        for name in [
            config.sheet_names.model1,
            config.sheet_names.model2,
            config.sheet_names.comparison,
        ]
    ]
    assert one['shared']['ID'].value == 'id-shared'
    assert two['shared']['ID'].value == 'prod-shared'
    assert one['shared']['ap50'].value == 'NA'
    assert one['new']['ap50'].value == 0
    assert one['Среднее']['ap50'].value == 0
    assert one['Среднее']['ap50 coverage'].value == '1/2'
    assert two['shared']['ap50'].value == 'NA'
    assert one['shared']['precision'].value == 'NA'
    assert two['shared']['precision'].value == 0.4
    assert diff['Среднее']['ap50'].value == 'NA'
    assert diff['Среднее']['ap50'].fill.patternType is None


@pytest.mark.parametrize('builder', [DevReportBuilder, BusinessReportBuilder])
@pytest.mark.parametrize('both_filtered', [False, True])
def test_empty_eligible_population(tmp_path, builder, both_filtered):
    config = Config.default()
    new = frame(['shared'], [0.9], [5])
    prod = frame(['shared'], [0.5], [5 if both_filtered else 20])
    paths = [tmp_path / 'new.xlsx', tmp_path / 'prod.xlsx']
    for data, path in zip([new, prod], paths):
        data.to_excel(path, index=False)
    output = tmp_path / 'out.xlsx'
    builder(*(MetricsReader(path) for path in paths), config).build(output)
    wb = load_workbook(output)
    one, diff = [
        table(wb[name]) for name in [config.sheet_names.model1, config.sheet_names.comparison]
    ]
    assert one['Среднее']['f1_score'].value == 'NA'
    assert one['Среднее']['f1_score coverage'].value == '0/0'
    assert ('shared' in one) is not both_filtered
    assert diff['Среднее']['f1_score'].value == 'NA'
    if builder is BusinessReportBuilder:
        verdict = wb[config.sheet_names.verdict]
        assert verdict['E6'].value == 'Недостаточно данных'
        for coordinate in ['B3', 'B4', 'B5', 'D3', 'D4', 'D5', 'E3', 'E4', 'E5']:
            assert verdict[coordinate].value == 'NA'
            assert verdict[coordinate].fill.patternType is None


def test_verdict_goal_population_excludes_display_placeholders(tmp_path):
    config = Config.default()
    new = frame(['shared', 'new'], [0.9, 0.9])
    new.loc[1, 'perebrak'] = 0.7
    prod = frame(['shared', 'deleted'], [0.5, 0.5])
    paths = [tmp_path / 'new.xlsx', tmp_path / 'prod.xlsx']
    for data, path in zip([new, prod], paths):
        data.to_excel(path, index=False)
    output = tmp_path / 'business.xlsx'
    BusinessReportBuilder(*(MetricsReader(path) for path in paths), config).build(output)
    wb = load_workbook(output)
    verdict = wb[config.sheet_names.verdict]
    assert verdict['B4'].value == 0.5 and verdict['C4'].value == 1
    assert verdict['B5'].value == 0.5 and verdict['C5'].value == 0
    assert verdict['E4'].value == verdict['E5'].value == 0
    assert verdict['E6'].value == 'Не к выкатке'
    one = table(wb[config.sheet_names.model1])
    assert one['Среднее'][config.business.col_perebrak_target].value == 0.5
    assert one['deleted'][config.business.col_perebrak_target].value == 'NA'


@pytest.mark.parametrize('source', ['f1_score', 'perebrak', 'nedobrak'])
@pytest.mark.parametrize('invalid', [-0.2, 1.2])
@pytest.mark.parametrize('custom_alias', [False, True])
def test_invalid_business_source_has_no_saved_numeric_difference(
    tmp_path, source, invalid, custom_alias
):
    config = Config.default()
    metric = source
    new, prod = frame(['shared'], [0.9]), frame(['shared'], [0.5])
    if custom_alias:
        metric = 'custom_business_metric'
        key = {'f1_score': 'f1_col', 'perebrak': 'perebrak_col', 'nedobrak': 'nedobrak_col'}[source]
        path = tmp_path / 'config.yaml'
        path.write_text(f'business:\n  {key}: {metric}\n', encoding='utf-8')
        config = Config.load(path)
        new = new.rename(columns={source: metric})
        prod = prod.rename(columns={source: metric})
    new.loc[0, metric] = invalid
    paths = [tmp_path / 'new.xlsx', tmp_path / 'prod.xlsx']
    for data, path in zip([new, prod], paths):
        data.to_excel(path, index=False)
    output = tmp_path / 'business.xlsx'
    with pytest.warns(UserWarning, match='Unavailable business values'):
        BusinessReportBuilder(*(MetricsReader(path) for path in paths), config).build(output)
    wb = load_workbook(output)
    display_metric = config.business.column_translations.get(metric, metric)
    for sheet_name in [config.sheet_names.model1, config.sheet_names.comparison]:
        rows = table(wb[sheet_name])
        for name in ['shared', 'Среднее']:
            assert rows[name][display_metric].value == 'NA'
            assert rows[name][display_metric].fill.patternType is None
    assert wb[config.sheet_names.verdict]['E6'].value == 'Недостаточно данных'
