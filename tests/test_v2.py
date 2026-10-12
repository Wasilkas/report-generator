"""Structured presentation preserves exact labels and legacy populations."""

from copy import copy

import pandas as pd
import pytest
from openpyxl import load_workbook

from report_generator.config import Config
from report_generator.core.reader import MetricsReader
from report_generator.reports.business.builder import BusinessReportBuilder as LegacyBusiness
from report_generator.reports.dev.builder import DevReportBuilder as LegacyDev
from report_generator.v2 import (
    BusinessReportBuilder,
    DevReportBuilder,
    MetricsTable,
    PresentationOptions,
    compact_report_counts,
    read_metrics_table,
)


def frame(labels, values, counts=None):
    return pd.DataFrame(
        {
            'Класс': labels,
            'Количество примеров train': counts or [30] * len(labels),
            'f1_score': values,
            'ap50': values,
            'perebrak': [0.1] * len(labels),
            'nedobrak': [0.2] * len(labels),
            'confidence': [0.5] * len(labels),
        }
    )


def sheet_snapshot(sheet):
    return {
        'cells': [
            (c.coordinate, c.value, c.data_type, c.number_format, str(c._style))
            for row in sheet
            for c in row
            if c.value is not None
        ],
        'freeze': sheet.freeze_panes,
        'widths': {key: value.width for key, value in sheet.column_dimensions.items()},
        'heights': {key: value.height for key, value in sheet.row_dimensions.items()},
        'merges': str(sheet.merged_cells),
    }


@pytest.mark.parametrize(
    'legacy,builder', [(LegacyDev, DevReportBuilder), (LegacyBusiness, BusinessReportBuilder)]
)
@pytest.mark.parametrize('scenario', ['shared', 'disjoint', 'asymmetric', 'empty'])
@pytest.mark.parametrize('counts_layout', ['columns', 'below_table'])
def test_structured_and_legacy_equivalent_calculations_and_layout(
    tmp_path,
    legacy,
    builder,
    scenario,
    counts_layout,
):
    if scenario == 'disjoint':
        one, two = frame(['new'], [0.0]), frame(['deleted'], [0.8])
    else:
        one = frame(
            ['shared', 'new', 'filtered'],
            [0.8, 0.0, None],
            [0, 0, 0]
            if scenario == 'empty'
            else ([30, 30, 0] if scenario == 'asymmetric' else [30, 30, 30]),
        )
        two = frame(
            ['shared', 'deleted', 'filtered'],
            [0.6, 0.4, 0.7],
            [0, 0, 0] if scenario == 'empty' else [30, 30, 30],
        )
    paths = [tmp_path / 'new.xlsx', tmp_path / 'prod.xlsx']
    for table, path in zip([one, two], paths):
        table.to_excel(path, index=False)
    config = Config.load()
    old_path, new_path = tmp_path / 'old.xlsx', tmp_path / 'new-report.xlsx'
    legacy(*(MetricsReader(path) for path in paths), config).build(old_path)
    if counts_layout == 'below_table':
        compact_report_counts(
            old_path,
            (config.sheet_names.model1, config.sheet_names.model2, config.sheet_names.comparison),
            translations=config.business.column_translations if legacy is LegacyBusiness else None,
        )
    tables = [read_metrics_table(path) for path in paths]
    before = [table.frame.copy(deep=True) for table in tables]
    builder(*tables, config, presentation=PresentationOptions(counts=counts_layout)).build(new_path)
    old, new = load_workbook(old_path), load_workbook(new_path)
    assert old.sheetnames == new.sheetnames
    for old_sheet, new_sheet in zip(old, new):
        # The direct layout registers styles only after removing columns; style
        # IDs are package-local, while the actual styles must remain equivalent.
        for a, b in zip(old_sheet, new_sheet):
            for c, d in zip(a, b):
                assert (
                    c.value,
                    c.data_type,
                    c.number_format,
                    copy(c.font),
                    copy(c.fill),
                    copy(c.alignment),
                ) == (
                    d.value,
                    d.data_type,
                    d.number_format,
                    copy(d.font),
                    copy(d.fill),
                    copy(d.alignment),
                )
        a, b = sheet_snapshot(old_sheet), sheet_snapshot(new_sheet)
        a.pop('cells')
        b.pop('cells')
        assert a == b
        assert old_sheet.max_row == new_sheet.max_row
        assert old_sheet.max_column == new_sheet.max_column
    for table, snapshot in zip(tables, before):
        pd.testing.assert_frame_equal(table.frame, snapshot)
    old.close()
    new.close()


@pytest.mark.parametrize('builder', [DevReportBuilder, BusinessReportBuilder])
def test_lexical_classes_numeric_footers_and_direct_availability(tmp_path, builder):
    labels = ['0001', '1', 'NA', ' 1 ', '0002', 0.5, 'Среднее', '=safe']
    source = tmp_path / 'source.xlsx'
    data = frame(
        labels, [0.1, 0.2, 0.3, None, 0.5, 0.99, 0.99, 0.2], [30, 30, 30, 30, 0, 30, 30, 30]
    )
    # The XLSX input needs literal string typing for formula-like class names.
    data.to_excel(source, index=False)
    workbook = load_workbook(source)
    workbook.active.cell(9, 1).data_type = 's'
    workbook.save(source)
    workbook.close()
    table = read_metrics_table(source)
    assert table.frame['Класс'].tolist() == ['0001', '1', 'NA', ' 1 ', '0002', '=safe']
    assert MetricsReader(source).read()['Класс'].tolist() == ['=safe']
    config = Config.load()
    output = tmp_path / 'output.xlsx'
    builder(table, table, config, presentation=PresentationOptions(counts='below_table')).build(
        output
    )
    workbook = load_workbook(output)
    for title in [
        config.sheet_names.model1,
        config.sheet_names.model2,
        config.sheet_names.comparison,
    ]:
        sheet = workbook[title]
        rows = {sheet.cell(r, 1).value: r for r in range(3, sheet.max_row + 1)}
        assert {'0001', '1', 'NA', ' 1 ', '=safe'} <= rows.keys()
        assert '0002' not in rows
        assert sheet.cell(rows['=safe'], 1).data_type == 's'
        assert sheet.cell(rows['AP50-valid-class-count'], 2).value == '4/5'
        assert not any(str(c.value).endswith(' coverage') for c in sheet[2])
    workbook.close()


@pytest.mark.parametrize('builder', [DevReportBuilder, BusinessReportBuilder])
def test_custom_translations_keep_metric_and_count_source_names(tmp_path, builder):
    settings = tmp_path / 'config.yaml'
    settings.write_text(
        'business:\n  column_translations:\n    ap50: AP custom\n'
        '    ap50 coverage: AP availability\n'
        '    absent coverage: Availability custom\n'
        '    unused: f1_score coverage\n'
    )
    data = frame(['A', 'B'], [0, None])
    data['Availability custom'] = [0.4, 0.6]
    table = MetricsTable(data, 'source.xlsx')
    config = Config.load(settings)
    output = tmp_path / 'result.xlsx'
    builder(table, table, config, presentation=PresentationOptions(counts='below_table')).build(
        output
    )
    workbook = load_workbook(output)
    for title in (
        config.sheet_names.model1,
        config.sheet_names.model2,
        config.sheet_names.comparison,
    ):
        sheet = workbook[title]
        headers = [c.value for c in sheet[2]]
        assert 'AP availability' not in headers
        assert 'Availability custom' in headers
        rows = {sheet.cell(r, 1).value: r for r in range(3, sheet.max_row + 1)}
        assert sheet.cell(rows['AP50-valid-class-count'], 2).value == '1/2'
        assert sheet.cell(rows['Availability custom-valid-class-count'], 2).value == '2/2'
    workbook.close()


@pytest.mark.parametrize(
    'labels,ids,error',
    [
        (['0001', '0001'], [1, 2], 'Duplicate classes'),
        (['0001', '1'], [1, 1], 'ID values'),
        (['0001', '1'], [1, None], 'ID values'),
        (['0001', '1'], [1, float('inf')], 'ID values'),
        (['0001', 0.5], [1, 2], 'lexical classes'),
    ],
)
def test_structured_validation(labels, ids, error):
    with pytest.raises(ValueError, match=error):
        MetricsTable(pd.DataFrame({'Класс': labels, 'ID': ids}), 'invalid.xlsx')


def test_structured_snapshot_and_output_overwrite_aliases(tmp_path):
    data = frame(['cat'], [0.8])
    table = MetricsTable(data, 'source.xlsx')
    data.loc[0, 'Класс'] = 'changed'
    assert table.frame.loc[0, 'Класс'] == 'cat'
    source = tmp_path / 'source.xlsx'
    data.to_excel(source, index=False)
    table = read_metrics_table(source)
    builder = DevReportBuilder(
        table, table, Config.load(), presentation=PresentationOptions(counts='below_table')
    )
    with pytest.raises(ValueError, match='alias'):
        builder.build(source, overwrite=True)
    for link in [tmp_path / 'symlink.xlsx', tmp_path / 'hardlink.xlsx']:
        if link.name.startswith('symlink'):
            link.symlink_to(source)
        else:
            link.hardlink_to(source)
        with pytest.raises(ValueError, match='alias'):
            builder.build(link, overwrite=True)
    output = tmp_path / 'output.xlsx'
    builder.build(output)
    with pytest.raises(FileExistsError):
        builder.build(output)
    builder.build(output, overwrite=True)


def test_explicit_count_title_and_labels_are_literal(tmp_path):
    table = MetricsTable(frame(['cat'], [0.8]), 'source.xlsx')
    options = PresentationOptions(
        counts='below_table', count_title='=literal', count_labels={'ap50': '=custom'}
    )
    output = tmp_path / 'output.xlsx'
    DevReportBuilder(table, table, Config.load(), presentation=options).build(output)
    workbook = load_workbook(output)
    for title in (
        Config.load().sheet_names.model1,
        Config.load().sheet_names.model2,
        Config.load().sheet_names.comparison,
    ):
        sheet = workbook[title]
        cells = {cell.value: cell for row in sheet for cell in row if isinstance(cell.value, str)}
        assert cells['=literal'].data_type == 's'
        assert cells['=custom-valid-class-count'].data_type == 's'
    workbook.close()
