"""ExcelSheetWriter: renders a DataFrame onto an openpyxl worksheet."""

from __future__ import annotations

import math
from typing import Any, Callable

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..config import Config
from .reader import CLASS_COL

# ── Static style constants (not configurable) ──────────────────────────────────

_HEADER_FONT = Font(bold=True, name='Calibri')
_MEAN_FONT = Font(bold=True, italic=True, name='Calibri')
_DEFAULT_FONT = Font(name='Calibri')

_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
_LEFT = Alignment(horizontal='left', vertical='center')

# ColorFn signature: (column_name, numeric_value) -> PatternFill | None
ColorFn = Callable[[str, float], 'PatternFill | None']


def _fill(hex_color: str) -> PatternFill:
    return PatternFill('solid', fgColor=hex_color)


# ── Writer ─────────────────────────────────────────────────────────────────────


class ExcelSheetWriter:
    """Writes a DataFrame (with optional 'Среднее' row) to a worksheet.

    Parameters
    ----------
    config:
        Report configuration (colors, column format sets, …).
    color_fn:
        Optional callable ``(col_name, value) -> PatternFill | None`` applied
        to every numeric data cell.
    format_map:
        Optional ``{column_name: excel_format_string}`` mapping.  When provided,
        it takes full priority over the config-based ratio_cols / int_cols logic
        and is applied to both int and float numeric values.  Columns absent from
        the map fall back to ``"0.0000"``.  When *not* provided, the existing
        config-based logic is used (floats only).
    """

    def __init__(
        self,
        config: Config,
        color_fn: ColorFn | None = None,
        format_map: dict[str, str] | None = None,
    ) -> None:
        self._config = config
        self._color_fn = color_fn
        self._format_map = format_map

        self._header_fill = _fill(config.colors.header)
        self._mean_fill = _fill(config.colors.mean)
        self._pos_fill = _fill(config.colors.positive)
        self._neg_fill = _fill(config.colors.negative)

    @classmethod
    def with_comparison_colors(cls, config: Config) -> 'ExcelSheetWriter':
        """Create a writer with per-column comparison-sheet coloring."""
        better_higher = config.better_higher_cols
        better_lower = config.better_lower_cols
        threshold = config.degradation_threshold
        pos = _fill(config.colors.positive)
        neg = _fill(config.colors.negative)

        def comparison_color(col: str, val: float) -> PatternFill | None:
            if col in better_higher:
                if val > 0:
                    return pos
                if val < -threshold:
                    return neg
            elif col in better_lower:
                if val < 0:
                    return pos
                if val > threshold:
                    return neg
            return None

        return cls(config, color_fn=comparison_color)

    # ── Public API ─────────────────────────────────────────────────────────────

    def write(self, ws: Worksheet, df: pd.DataFrame, sheet_title: str = '') -> None:
        row_offset = 0

        if sheet_title:
            title_cell = ws.cell(1, 1, sheet_title)
            title_cell.font = Font(bold=True, size=13, name='Calibri')
            row_offset = 1

        self._write_header(ws, df, row_offset)
        self._write_rows(ws, df, row_offset)
        self._autofit_columns(ws)
        self._freeze_header(ws, row_offset)

    # ── Private helpers ────────────────────────────────────────────────────────

    def _write_header(self, ws: Worksheet, df: pd.DataFrame, row_offset: int) -> None:
        for c_idx, col in enumerate(df.columns, 1):
            cell = ws.cell(row_offset + 1, c_idx)
            cell.value = None if col == CLASS_COL else col
            cell.font = _HEADER_FONT
            cell.fill = self._header_fill
            cell.alignment = _CENTER

    def _write_rows(self, ws: Worksheet, df: pd.DataFrame, row_offset: int) -> None:
        col_names = list(df.columns)
        for r_idx, (_, row) in enumerate(df.iterrows(), row_offset + 2):
            is_mean = str(row.get(CLASS_COL, '')) == 'Среднее'

            for c_idx, col in enumerate(col_names, 1):
                raw = row[col]
                val = _clean(raw)
                cell = ws.cell(r_idx, c_idx)
                cell.value = val
                cell.font = _MEAN_FONT if is_mean else _DEFAULT_FONT
                cell.alignment = _LEFT if col == CLASS_COL else _CENTER

                # Fill
                if is_mean:
                    cell.fill = self._mean_fill
                elif self._color_fn is not None and isinstance(val, (int, float)):
                    fill = self._color_fn(col, val)
                    if fill is not None:
                        cell.fill = fill

                # Number format
                if isinstance(val, (int, float)) and not (
                    isinstance(val, float) and math.isnan(val)
                ):
                    if self._format_map is not None:
                        cell.number_format = self._format_map.get(col, '0.0000')
                    elif isinstance(val, float):
                        if col in self._config.ratio_cols:
                            cell.number_format = '0.0000'
                        elif col in self._config.int_cols:
                            cell.number_format = '0'
                        else:
                            cell.number_format = '0.0000'

    def _autofit_columns(self, ws: Worksheet) -> None:
        for col_cells in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col_cells[0].column)
            for cell in col_cells:
                try:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 40)

    def _freeze_header(self, ws: Worksheet, row_offset: int) -> None:
        ws.freeze_panes = ws.cell(row_offset + 2, 2)


# ── Helpers ────────────────────────────────────────────────────────────────────


def _clean(val: Any) -> Any:
    """Convert NaN / None to None so openpyxl stores a blank cell."""
    if val is None:
        return None
    try:
        if pd.isna(val):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(val, float) and math.isnan(val):
        return None
    return val
