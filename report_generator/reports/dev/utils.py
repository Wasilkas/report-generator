"""Utility functions for the developer Excel report."""

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ...app_config import CLASS_COL, REASON_COL, TRAIN_COUNT_DISPLAY_COL
from ...config import Config


def write_excluded_sheet(ws: Worksheet, excluded: pd.DataFrame, config: Config) -> None:
    header_font = Font(bold=True, name='Calibri')
    header_fill = PatternFill('solid', fgColor=config.colors.header)
    center = Alignment(horizontal='center', vertical='center')

    ws.cell(
        1, 1, 'Исключенные из попарного сравнения классы — причины и исходные модели'
    ).font = Font(bold=True, size=13, name='Calibri')

    display = {CLASS_COL: 'Класс'}
    for c_idx, col in enumerate(excluded.columns, 1):
        cell = ws.cell(2, c_idx, display.get(col, col))
        cell.data_type = 's'
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for r, (_, row) in enumerate(excluded.iterrows(), 3):
        for c_idx, col in enumerate(excluded.columns, 1):
            val = row[col]
            cell = ws.cell(r, c_idx, None if pd.isna(val) else val)
            if isinstance(val, str):
                cell.data_type = 's'
            cell.font = Font(name='Calibri')
            cell.alignment = center
            if isinstance(val, (int, float)):
                cell.number_format = '0'

    col_widths = {CLASS_COL: 30, TRAIN_COUNT_DISPLAY_COL: 22, REASON_COL: 45}
    for c_idx, col in enumerate(excluded.columns, 1):
        ws.column_dimensions[get_column_letter(c_idx)].width = col_widths.get(col, 20)

    ws.freeze_panes = ws.cell(3, 1)
