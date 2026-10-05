"""DevReportBuilder: assembles the four-sheet developer report."""

from pathlib import Path

from openpyxl import Workbook

from ...config import Config
from ...core.reader import MetricsReader
from ...core.writer import ExcelSheetWriter, write_comparison_legend
from ..base import BaseReportBuilder
from .utils import write_excluded_sheet


class DevReportBuilder(BaseReportBuilder):
    """Builds a four-sheet Excel developer report.

    Sheet 1 — model1 (fresh / candidate) metrics + 'Среднее' row.
    Sheet 2 — model2 (production) metrics + 'Среднее' row.
    Sheet 3 — numeric difference (model1 − model2) + 'Среднее' row of diffs.
    Sheet 4 — classes excluded due to insufficient training examples.

    All behaviour is controlled by *config*.
    """

    def __init__(
        self,
        model1_reader: MetricsReader,
        model2_reader: MetricsReader,
        config: Config,
    ) -> None:
        super().__init__(model1_reader, model2_reader, config)
        self._sheet_writer = ExcelSheetWriter(config)
        self._diff_writer = ExcelSheetWriter.with_comparison_colors(config)

    def build(self, output_path: str | Path, *, overwrite: bool = False) -> None:
        output_path = Path(output_path)
        self._check_output(output_path, overwrite)
        names = self._config.sheet_names

        df1_raw = self.model1_reader.read()
        df2_raw = self.model2_reader.read()

        # Assign global IDs before filtering so excluded classes leave gaps.
        for raw in (df1_raw, df2_raw):
            if 'ID' not in raw.columns:
                raw.insert(1, 'ID', range(len(raw)))

        df1, df2, excluded = self._prepare(df1_raw, df2_raw)

        # Comparison: drop ID to avoid it appearing as a numeric diff.
        df_diff = self._cmp_calc.compute(df1.drop(columns=['ID']), df2.drop(columns=['ID']))

        df1_with_mean = self._mean_calc.append_mean_row(df1)
        df2_with_mean = self._mean_calc.append_mean_row(df2)
        df_diff_with_mean = self._mean_calc.append_mean_row(df_diff)

        wb = Workbook()
        ws1 = wb.active
        ws1.title = names.model1
        ws2 = wb.create_sheet(names.model2)
        ws3 = wb.create_sheet(names.comparison)
        ws4 = wb.create_sheet(names.excluded)

        self._sheet_writer.write(
            ws1,
            df1_with_mean,
            sheet_title=f'Метрики новой модели — {self.model1_reader.file_path.name}',
        )
        self._sheet_writer.write(
            ws2,
            df2_with_mean,
            sheet_title=f'Метрики прод модели — {self.model2_reader.file_path.name}',
        )
        last_row = self._diff_writer.write(
            ws3,
            df_diff_with_mean,
            sheet_title='Сравнение (новая − прод)',
        )
        write_comparison_legend(ws3, last_row + 2, self._config)
        write_excluded_sheet(ws4, excluded, self._config)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        print(f'Dev report saved → {output_path}')
