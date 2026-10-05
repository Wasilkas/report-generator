"""Utility functions for the business Excel report."""

import math
import warnings

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ...app_config import CLASS_COL, REASON_COL, TRAIN_COUNT_DISPLAY_COL
from ...config import BusinessConfig, Config
from ...core.writer import (
    _CENTER,
    _DEFAULT_FONT,
    _HEADER_FONT,
    _LEFT,
    _MEAN_FONT,
    ColorFn,
    _fill,
)

# ── Data preparation ──────────────────────────────────────────────────────────


def add_goal_columns(df: pd.DataFrame, biz: BusinessConfig) -> pd.DataFrame:
    """Append four binary (0/1) goal columns after the nedobrak column."""
    df = df.copy()

    collisions = set(biz.goal_cols) & set(df.columns)
    if collisions:
        raise ValueError(f'Generated goal column collision: {sorted(collisions)}')
    # Business criteria require finite ratios. AP NaNs stay untouched.
    for source in [biz.f1_col, biz.perebrak_col, biz.nedobrak_col]:
        raw = df[source] if source in df else pd.Series(float('nan'), index=df.index)
        values = pd.to_numeric(raw, errors='coerce')
        valid = values.map(
            lambda value: pd.notna(value) and math.isfinite(value) and 0 <= value <= 1
        ) & ~raw.map(lambda value: isinstance(value, bool))
        if not valid.all() or source not in df:
            warnings.warn(
                f'Unavailable business values in {source!r}: '
                f'{int((~valid).sum())}/{len(df)} classes; rollout requires all criteria',
                UserWarning,
                stacklevel=2,
            )
        df[source] = values.where(valid)
    for source, target, goal, gross in [
        (biz.perebrak_col, biz.target_perebrak, biz.col_perebrak_target, biz.col_perebrak_gross),
        (biz.nedobrak_col, biz.target_nedobrak, biz.col_nedobrak_target, biz.col_nedobrak_gross),
    ]:
        values = df[source]
        df[goal] = (values <= target).astype(float).where(values.notna())
        df[gross] = (values > 2 * target).astype(float).where(values.notna())

    # Reorder: insert goal columns right after nedobrak
    goal_cols = biz.goal_cols
    cols = list(df.columns)
    for g in goal_cols:
        if g in cols:
            cols.remove(g)
    insert_after = biz.nedobrak_col if biz.nedobrak_col in cols else (cols[-1] if cols else None)
    if insert_after is not None:
        pos = cols.index(insert_after) + 1
        for g in reversed(goal_cols):
            cols.insert(pos, g)
    return df[cols]


def to_percentage(df: pd.DataFrame, biz: BusinessConfig) -> pd.DataFrame:
    """Multiply perebrak and nedobrak columns by 100 (copy)."""
    df = df.copy()
    for col in [biz.perebrak_col, biz.nedobrak_col]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors='coerce') * 100
    return df


def translate_columns(df: pd.DataFrame, translations: dict[str, str]) -> pd.DataFrame:
    rename = {col: translations.get(col, col) for col in df.columns}
    if len(set(rename.values())) != len(rename):
        raise ValueError('Translated column name collision')
    return df.rename(columns=rename)


# ── Format map & comparison color fn ─────────────────────────────────────────


def build_format_map(config: Config) -> dict[str, str]:
    """Build column → Excel format string map for the business sheets."""
    biz = config.business
    trans = biz.column_translations
    fmt: dict[str, str] = {}

    # Translated percentage columns
    for src in [biz.perebrak_col, biz.nedobrak_col]:
        fmt[trans.get(src, src)] = '0.00'

    # Ratio columns (names unchanged by translation)
    for col in config.ratio_cols:
        fmt[trans.get(col, col)] = '0.0000'

    for src in [biz.perebrak_col, biz.nedobrak_col]:
        fmt[trans.get(src, src)] = '0.00'

    # Integer columns (use translated names)
    for col in config.int_cols:
        fmt[trans.get(col, col)] = '0'

    # Goal columns — 2 dp so mean row (fraction) displays nicely
    for col in biz.goal_cols:
        fmt[trans.get(col, col)] = '0.00'

    return fmt


def build_comparison_color_fn(config: Config) -> ColorFn:
    """Build comparison color fn for the business Сравнение sheet."""
    biz = config.business
    trans = biz.column_translations
    pos = _fill(config.colors.positive)
    neg = _fill(config.colors.negative)

    better_higher = {trans.get(col, col) for col in config.better_higher_cols | {biz.f1_col}}
    dth = config.degradation_threshold

    pct_cols = {trans.get(col, col) for col in [biz.perebrak_col, biz.nedobrak_col]}
    pth = biz.comparison_pct_threshold

    def color_fn(col: str, val: float) -> PatternFill | None:
        if col in better_higher:
            if val > 0:
                return pos
            if val < -dth:
                return neg
        elif col in pct_cols:
            if val < 0:
                return pos
            if val > pth:
                return neg
        return None

    return color_fn


# ── Excluded-classes sheet ────────────────────────────────────────────────────


def write_excluded_sheet(ws: Worksheet, excluded: pd.DataFrame, config: Config) -> None:
    header_fill = _fill(config.colors.header)
    trans = config.business.column_translations

    ws.cell(
        1, 1, 'Исключенные из попарного сравнения классы — причины и исходные модели'
    ).font = Font(bold=True, size=13, name='Calibri')

    for c_idx, col in enumerate(excluded.columns, 1):
        display = trans.get(col, col)
        cell = ws.cell(2, c_idx, display)
        cell.data_type = 's'
        cell.font = _HEADER_FONT
        cell.fill = header_fill
        cell.alignment = _CENTER

    for r, (_, row) in enumerate(excluded.iterrows(), 3):
        for c_idx, col in enumerate(excluded.columns, 1):
            val = row[col]
            cell = ws.cell(r, c_idx, None if pd.isna(val) else val)
            if isinstance(val, str):
                cell.data_type = 's'
            cell.font = _DEFAULT_FONT
            cell.alignment = _CENTER
            if isinstance(val, (int, float)):
                cell.number_format = '0'

    col_widths = {CLASS_COL: 30, TRAIN_COUNT_DISPLAY_COL: 22, REASON_COL: 45}
    for c_idx, col in enumerate(excluded.columns, 1):
        ws.column_dimensions[get_column_letter(c_idx)].width = col_widths.get(col, 20)

    ws.freeze_panes = ws.cell(3, 1)


# ── Verdict computation ───────────────────────────────────────────────────────


def f1_mean(df: pd.DataFrame, biz: BusinessConfig) -> float | None:
    if biz.f1_col not in df.columns:
        return None
    vals = pd.to_numeric(df[biz.f1_col], errors='coerce')
    if (
        len(vals) == 0
        or not vals.map(lambda v: pd.notna(v) and math.isfinite(v) and 0 <= v <= 1).all()
    ):
        return None
    if df[biz.f1_col].map(lambda value: isinstance(value, bool)).any():
        return None
    return float(vals.mean())


def classes_in_target(df: pd.DataFrame, biz: BusinessConfig) -> float | None:
    if biz.col_perebrak_target not in df.columns or biz.col_nedobrak_target not in df.columns:
        return None
    p = pd.to_numeric(df[biz.col_perebrak_target], errors='coerce')
    n = pd.to_numeric(df[biz.col_nedobrak_target], errors='coerce')
    if not p.isin([0, 1]).all() or not n.isin([0, 1]).all():
        return None
    both = ((p == 1) & (n == 1)).sum()
    return float(both / len(df)) if len(df) > 0 else None


def classes_gross_not_target(df: pd.DataFrame, biz: BusinessConfig) -> float | None:
    if biz.col_perebrak_gross not in df.columns or biz.col_nedobrak_gross not in df.columns:
        return None
    p = pd.to_numeric(df[biz.col_perebrak_gross], errors='coerce')
    n = pd.to_numeric(df[biz.col_nedobrak_gross], errors='coerce')
    if not p.isin([0, 1]).all() or not n.isin([0, 1]).all():
        return None
    both = ((p == 1) | (n == 1)).sum()
    return float(both / len(df)) if len(df) > 0 else None


def rel_diff(new_val: float | None, prod_val: float | None) -> float | None:
    if (
        new_val is None
        or prod_val is None
        or not math.isfinite(new_val)
        or not math.isfinite(prod_val)
    ):
        return None
    if prod_val == 0:
        return None
    return (new_val - prod_val) / abs(prod_val)


def score_higher_better(rel: float | None, threshold: float) -> int | None:
    if rel is None:
        return None
    if rel == 0:
        return 1
    if rel < 0:
        return 0
    if rel < threshold:
        return 1
    return 2


def score_lower_better(rel: float | None, threshold: float) -> int | None:
    if rel is None:
        return None
    if rel == 0:
        return 1
    if rel > 0:
        return 0
    if rel > -threshold:
        return 1
    return 2


# ── Verdict sheet writer ──────────────────────────────────────────────────────


def write_verdict_sheet(
    ws: Worksheet,
    df1_goals: pd.DataFrame,
    df2_goals: pd.DataFrame,
    config: Config,
) -> None:
    biz = config.business
    threshold = biz.verdict_score_threshold

    criteria = [
        ('F1 mean', f1_mean(df1_goals, biz), f1_mean(df2_goals, biz), 'higher'),
        (
            'Классы в целях',
            classes_in_target(df1_goals, biz),
            classes_in_target(df2_goals, biz),
            'higher',
        ),
        (
            'Классы грубо не в целях',
            classes_gross_not_target(df1_goals, biz),
            classes_gross_not_target(df2_goals, biz),
            'lower',
        ),
    ]

    header_fill = _fill(config.colors.header)
    mean_fill = _fill(config.colors.mean)
    pos_fill = _fill(config.colors.positive)
    neg_fill = _fill(config.colors.negative)

    ws.cell(1, 1, 'Вердикт').font = Font(bold=True, size=13, name='Calibri')

    for c_idx, h in enumerate(
        ['Критерий', 'Новая модель', 'Прод модель', 'Относительная разница, %', 'Итого'], 1
    ):
        cell = ws.cell(2, c_idx, h)
        cell.data_type = 's'
        cell.font = _HEADER_FONT
        cell.fill = header_fill
        cell.alignment = _CENTER

    scores: list[int] = []
    for r_idx, (name, new_val, prod_val, direction) in enumerate(criteria, 3):
        r = rel_diff(new_val, prod_val)
        rel_pct = r * 100 if r is not None else None
        score = (
            score_higher_better(r, threshold)
            if direction == 'higher'
            else score_lower_better(r, threshold)
        )
        if new_val is not None and prod_val == 0 and math.isfinite(new_val):
            score = 1 if new_val == 0 else (2 if direction == 'higher' else 0)
        if score is not None:
            scores.append(score)

        for c_idx, val in enumerate([name, new_val, prod_val, rel_pct, score], 1):
            cell = ws.cell(r_idx, c_idx, 'NA' if val is None else val)
            if isinstance(val, str):
                cell.data_type = 's'
            cell.font = _DEFAULT_FONT
            cell.alignment = _CENTER if c_idx > 1 else _LEFT
            if isinstance(val, float) and not math.isnan(val):
                cell.number_format = '0.00' if c_idx == 4 else '0.0000'
            elif isinstance(val, int):
                cell.number_format = '0'
            if c_idx == 5 and score is not None:
                if score == 2:
                    cell.fill = pos_fill
                elif score == 0:
                    cell.fill = neg_fill

    if len(scores) != len(criteria):
        verdict_text, verdict_fill = 'Недостаточно данных', neg_fill
    elif any(s == 2 for s in scores) and all(s != 0 for s in scores):
        verdict_text, verdict_fill = 'К выкатке', pos_fill
    else:
        verdict_text, verdict_fill = 'Не к выкатке', neg_fill

    total_row = len(criteria) + 3
    for c_idx in range(1, 6):
        cell = ws.cell(total_row, c_idx)
        cell.font = _MEAN_FONT
        cell.fill = mean_fill
        cell.alignment = _LEFT if c_idx == 1 else _CENTER

    ws.cell(total_row, 1, 'Итого').font = _MEAN_FONT
    verdict_cell = ws.cell(total_row, 5, verdict_text)
    verdict_cell.font = _MEAN_FONT
    verdict_cell.fill = verdict_fill
    verdict_cell.alignment = _CENTER

    ws.column_dimensions['A'].width = 30
    for col_letter in ['B', 'C', 'D', 'E']:
        ws.column_dimensions[col_letter].width = 22
    ws.freeze_panes = ws.cell(3, 1)

    legend_start = total_row + 2
    label = ws.cell(legend_start, 1, 'Легенда (Итого):')
    label.font = Font(bold=True, name='Calibri')
    legend_entries = [
        (neg_fill, '0 — Хуже'),
        (None, '1 — Без изменений или незначительно лучше'),
        (pos_fill, '2 — Лучше'),
    ]
    for i, (fill, text) in enumerate(legend_entries, legend_start + 1):
        cell = ws.cell(i, 1, text)
        cell.font = Font(name='Calibri')
        cell.alignment = _LEFT
        if fill is not None:
            cell.fill = fill
