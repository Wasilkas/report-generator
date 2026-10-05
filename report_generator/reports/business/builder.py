"""BusinessReportBuilder: assembles the business-oriented Excel report."""

from pathlib import Path

from openpyxl import Workbook

from ...app_config import CLASS_COL
from ...config import Config
from ...core.reader import MetricsReader
from ...core.writer import ExcelSheetWriter, write_comparison_legend
from ..base import BaseReportBuilder
from .utils import (
    add_goal_columns,
    build_comparison_color_fn,
    build_format_map,
    to_percentage,
    translate_columns,
    write_excluded_sheet,
    write_verdict_sheet,
)


class BusinessReportBuilder(BaseReportBuilder):
    """Builds the business-oriented Excel report.

    Sheet 1 — new model metrics (Russian headers, % perebrak/nedobrak, goal columns).
    Sheet 2 — prod model metrics, same structure.
    Sheet 3 — Сравнение: numeric diff (new − prod) with comparison coloring.
    Sheet 4 — Вердикт: F1/classes-in-target/gross-not-target verdict table.
    Sheet 5 — Удаленные классы: classes excluded due to insufficient training data.
    """

    def __init__(
        self,
        model1_reader: MetricsReader,
        model2_reader: MetricsReader,
        config: Config,
    ) -> None:
        super().__init__(model1_reader, model2_reader, config)

        fmt = build_format_map(config)
        self._sheet_writer = ExcelSheetWriter(
            config,
            format_map=fmt,
            class_col=config.business.column_translations.get('Класс', 'Класс'),
        )
        self._diff_writer = ExcelSheetWriter(
            config,
            color_fn=build_comparison_color_fn(config),
            format_map=fmt,
            class_col=config.business.column_translations.get('Класс', 'Класс'),
        )

    def build(self, output_path: str | Path, *, overwrite: bool = False) -> None:
        output_path = Path(output_path)
        self._check_output(output_path, overwrite)
        biz = self._config.business
        translations = dict(biz.column_translations)

        df1_raw = self.model1_reader.read()
        df2_raw = self.model2_reader.read()

        df1, df2, excluded = self._prepare(df1_raw, df2_raw)

        df1_goals = add_goal_columns(df1, biz)
        df2_goals = add_goal_columns(df2, biz)

        classes = sorted(set(df1[CLASS_COL]) | set(df2[CLASS_COL]))
        population_note = self._population_note(df1, df2)
        # Verdict inputs remain unpadded; placeholders must not alter denominators.
        df_diff_display = translate_columns(
            to_percentage(
                self._comparison_with_mean(
                    df1_goals.drop(columns=biz.goal_cols),
                    df2_goals.drop(columns=biz.goal_cols),
                    classes,
                ),
                biz,
            ),
            translations,
        )
        df1_display = translate_columns(
            to_percentage(self._display_with_mean(df1_goals, classes), biz), translations
        )
        df2_display = translate_columns(
            to_percentage(self._display_with_mean(df2_goals, classes), biz), translations
        )

        wb = Workbook()
        ws1 = wb.active
        ws1.title = self._config.sheet_names.model1
        ws2 = wb.create_sheet(self._config.sheet_names.model2)
        ws3 = wb.create_sheet(self._config.sheet_names.comparison)
        ws_verdict = wb.create_sheet(self._config.sheet_names.verdict)
        ws_excl = wb.create_sheet(self._config.sheet_names.excluded)

        self._sheet_writer.write(
            ws1,
            df1_display,
            sheet_title=f'Метрики новой модели — {self.model1_reader.file_path.name}',
        )
        self._sheet_writer.write(
            ws2,
            df2_display,
            sheet_title=f'Метрики прод модели — {self.model2_reader.file_path.name}',
        )
        last_row = self._diff_writer.write(
            ws3, df_diff_display, sheet_title='Сравнение (новая − прод)'
        )
        write_comparison_legend(ws3, last_row + 2, self._config, business=True)
        write_verdict_sheet(ws_verdict, df1_goals, df2_goals, self._config)
        write_excluded_sheet(ws_excl, excluded, self._config)
        for ws in (ws1, ws2, ws3, ws_verdict):
            ws.cell(ws.max_row + 2, 1, population_note)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        print(f'Business report saved → {output_path}')
