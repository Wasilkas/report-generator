"""Additive structured report entry points with explicit layout selection."""

import re
from collections.abc import Mapping
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

from openpyxl.styles import Alignment, Font

from ..config import Config
from ..reports.business.builder import BusinessReportBuilder as LegacyBusinessReportBuilder
from ..reports.dev.builder import DevReportBuilder as LegacyDevReportBuilder
from .metrics import MetricsTable


@dataclass(frozen=True)
class PresentationOptions:
    """Choose availability placement without changing aggregation populations."""

    counts: Literal['below_table', 'columns']
    count_title: str = 'Valid class counts (valid / eligible)'
    count_labels: Mapping[str, str] = field(
        default_factory=lambda: {
            'ap50': 'AP50',
            'ap75': 'AP75',
            'ap50_95': 'AP50-95',
        }
    )

    def __post_init__(self) -> None:
        if self.counts not in ('below_table', 'columns'):
            raise ValueError(f'Unknown count presentation: {self.counts!r}')
        if not isinstance(self.count_title, str) or not self.count_title:
            raise ValueError('Count title must be nonempty text')
        if any(
            not isinstance(key, str) or not isinstance(value, str)
            for key, value in self.count_labels.items()
        ):
            raise ValueError('Count labels must map text metrics to text labels')
        object.__setattr__(self, 'count_labels', dict(self.count_labels))


class _TableReader:
    def __init__(self, table: MetricsTable) -> None:
        self.table = table
        self.file_path = table.source_path or Path(table.source_name)

    def read(self):
        return self.table.frame.copy(deep=True)


class _StructuredPresentation:
    def _init_tables(
        self, model1: MetricsTable, model2: MetricsTable, presentation: PresentationOptions
    ) -> None:
        self._tables = (model1, model2)
        self._presentation = presentation
        self._pending_counts = {}

    def _check_output(self, output_path: Path, overwrite: bool) -> None:
        for table in self._tables:
            source = table.source_path
            if source is not None and (
                output_path.resolve() == source.resolve()
                or (output_path.exists() and source.exists() and output_path.samefile(source))
            ):
                raise ValueError(f'Output must not alias an input workbook: {output_path}')
        if output_path.exists() and not overwrite:
            raise FileExistsError(
                f'Output already exists: {output_path}; use --force to replace it'
            )

    def _write_metrics(self, writer, sheet, frame, *, sheet_title):
        if self._presentation.counts == 'columns':
            return writer.write(sheet, frame, sheet_title=sheet_title)
        # The calculator appends generated coverage columns. Identify the
        # aggregate values directly, before rendering, so translated user metric
        # headings which happen to end in "coverage" remain ordinary metrics.
        class_col = writer._class_col
        mean = frame.loc[frame[class_col] == 'Среднее']
        if len(mean) != 1:
            raise ValueError(f'Report {sheet.title!r} must have exactly one mean row')
        translations = self._config.business.column_translations if self._business else {}
        reverse = {
            display: source
            for source, display in translations.items()
            if source.endswith(' coverage')
            and translations.get(source.removesuffix(' coverage'), source.removesuffix(' coverage'))
            in frame.columns
        }
        columns = []
        counts = []
        for column in frame.columns:
            value = mean.iloc[0][column]
            if not isinstance(value, str) or re.fullmatch(r'\d+/\d+', value) is None:
                continue
            source = reverse.get(column, column)
            if not source.endswith(' coverage'):
                raise ValueError(
                    f'Report {sheet.title!r} has an unknown class-count header {source!r}'
                )
            metric = source.removesuffix(' coverage')
            label = self._presentation.count_labels.get(metric, metric)
            counts.append((f'{label}-valid-class-count', value))
            columns.append(column)
        if columns and list(frame.columns[-len(columns) :]) != columns:
            raise ValueError(f'Report {sheet.title!r} has non-trailing coverage columns')
        self._pending_counts[sheet.title] = counts
        return writer.write(sheet, frame.drop(columns=columns), sheet_title=sheet_title)

    def _finish_metrics_sheet(self, sheet):
        counts = self._pending_counts.pop(sheet.title, [])
        if not counts:
            return
        start = sheet.max_row + 2
        title = sheet.cell(start, 1, self._presentation.count_title)
        title.data_type = 's'
        title.font = Font(name='Calibri', bold=True)
        title.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[start].height = 32
        for row, (label, value) in enumerate(counts, start + 1):
            name = sheet.cell(row, 1, label)
            name.data_type = 's'
            name.alignment = Alignment(wrap_text=True, vertical='center')
            name.font = Font(name='Calibri', bold=True)
            count = sheet.cell(row, 2, value)
            count.data_type = 's'
            count.alignment = Alignment(horizontal='center', vertical='center')
            sheet.row_dimensions[row].height = 30
        sheet.column_dimensions['A'].width = max(sheet.column_dimensions['A'].width, 36)


class DevReportBuilder(_StructuredPresentation, LegacyDevReportBuilder):
    """Developer reporting over supplied exact-label tables; no file rereads."""

    _business = False

    def __init__(
        self,
        model1: MetricsTable,
        model2: MetricsTable,
        config: Config,
        *,
        presentation: PresentationOptions,
    ) -> None:
        self._init_tables(model1, model2, presentation)
        super().__init__(_TableReader(model1), _TableReader(model2), config)


class BusinessReportBuilder(_StructuredPresentation, LegacyBusinessReportBuilder):
    """Business reporting retaining legacy calculation/verdict semantics."""

    _business = True

    def __init__(
        self,
        model1: MetricsTable,
        model2: MetricsTable,
        config: Config,
        *,
        presentation: PresentationOptions,
    ) -> None:
        self._init_tables(model1, model2, presentation)
        super().__init__(_TableReader(model1), _TableReader(model2), config)
